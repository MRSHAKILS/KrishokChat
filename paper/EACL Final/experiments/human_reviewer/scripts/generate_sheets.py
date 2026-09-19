import os
import json
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

BASE_DIR = r"D:\KrishokChat Advisory System\paper\EACL Final\experiments\human_reviewer"
SOURCE_JSON = r"D:\KrishokChat Advisory System\paper\EACL Final\experiments\N01_blind_arm\labeling_sheet_200.json"

os.makedirs(os.path.join(BASE_DIR, "batches"), exist_ok=True)
os.makedirs(os.path.join(BASE_DIR, "scripts"), exist_ok=True)
os.makedirs(os.path.join(BASE_DIR, "reference_ground_truth"), exist_ok=True)

# Load 200 queries
with open(SOURCE_JSON, "r", encoding="utf-8") as f:
    queries = json.load(f)

# Controlled crop list
CROPS_TAXONOMY = [
    ("none", "কোনো ফসল নেই / অনুপস্থিত", "Missing / Underspecified", "গাছের, পাতার, জমিতে, ইত্যাদি"),
    ("rice", "ধান", "Poaceae", "আমন, বোরো, আউশ, ব্রি ধান"),
    ("potato", "আলু", "Solanaceae", "গোল আলু, ডায়মন্ড, কার্ডিনাল"),
    ("wheat", "গম", "Poaceae", "বারি গম, গম ক্ষেত"),
    ("corn", "ভুট্টা", "Poaceae", "ভুট্টা, হাইব্রিড ভুট্টা"),
    ("tomato", "টমেটো", "Solanaceae", "টমেটো, বিলাতি বেগুন"),
    ("chilli", "মরিচ", "Solanaceae", "মরিচ, কাঁচা মরিচ, লঙ্কা"),
    ("brinjal", "বেগুন", "Solanaceae", "বেগুন, তাল বেগুন"),
    ("mustard", "সরিষা", "Brassicaceae", "সরিষা, টরি, রাই"),
    ("mango", "আম", "Anacardiaceae", "আম, হিমসাগর, ল্যাংড়া, ফজলি"),
    ("papaya", "পেঁপে", "Caricaceae", "পেঁপে, পেপে"),
    ("banana", "কলা", "Musaceae", "কলা, সাগর কলা, সবরি কলা"),
    ("guava", "পেয়ারা", "Myrtaceae", "পেয়ারা, কাজী পেয়ারা"),
    ("lemon", "লেবু", "Rutaceae", "লেবু, কাগজি লেবু, বাতাবি লেবু"),
    ("watermelon", "তরমুজ", "Cucurbitaceae", "তরমুজ, তরমুজের লতা"),
    ("coconut", "নারিকেল", "Arecaceae", "নারিকেল, নারিকেল গাছ, ডাব"),
    ("jackfruit", "কাঁঠাল", "Moraceae", "কাঁঠাল, কাঁঠাল গাছ"),
    ("cucumber", "শসা", "Cucurbitaceae", "শসা, শশা"),
    ("bottle_gourd", "লাউ", "Cucurbitaceae", "লাউ, কদু"),
    ("bitter_gourd", "করলা", "Cucurbitaceae", "করলা, উচ্ছে"),
    ("cabbage", "বাঁধাকপি", "Brassicaceae", "বাঁধাকপি, পাতাকপি"),
    ("cauliflower", "ফুলকপি", "Brassicaceae", "ফুলকপি"),
    ("jute", "পাট", "Malvaceae", "পাট, তোষা পাট, দেশী পাট"),
    ("onion", "পেঁয়াজ", "Amaryllidaceae", "পেঁয়াজ, পেয়াজ"),
    ("garlic", "রসুন", "Amaryllidaceae", "রসুন"),
    ("ginger", "আদা", "Zingiberaceae", "আদা"),
    ("turmeric", "হলুদ", "Zingiberaceae", "হলুদ"),
    ("betel_leaf", "পান", "Piperaceae", "পান, পানের বরজ"),
    ("tea", "চা", "Theaceae", "চা, চা বাগান"),
    ("mushroom", "মাশরুম", "Fungi", "মাশরুম"),
    ("strawberry", "স্ট্রবেরি", "Rosaceae", "স্ট্রবেরি"),
    ("dragon_fruit", "ড্রাগন ফল", "Cactaceae", "ড্রাগন ফল, ড্রাগন গাছ"),
    ("litchi", "লিচু", "Sapindaceae", "লিচু, বোম্বাই লিচু"),
    ("pointed_gourd", "পটল", "Cucurbitaceae", "পটল"),
    ("okra", "ঢেঁড়শ", "Malvaceae", "ঢেঁড়শ, ভেন্ডি"),
    ("bean", "শিম", "Fabaceae", "শিম, সীম"),
    ("other_crop", "অন্যান্য ফসল", "Other", "অন্য কোনো নির্দিষ্ট ফসল")
]

INTENT_OPTIONS = [
    ("Treatment", "রোগ বা পোকার প্রতিকার", "কীটনাশক, বালাইনাশক, রোগ লক্ষণ, স্প্রে, ওষুধ, বিষ, পোকা দমন"),
    ("Fertilizer", "সার ব্যবস্থাপনা ও মাত্রা", "ইউরিয়া, টিএসপি, ডিএপি, পটাশ, জিংক, সারের মাত্রা, প্রয়োগ সময়"),
    ("Prevention", "রোগ প্রতিরোধ ও শোধন", "বীজ শোধন, বালাই প্রতিরোধ, পূর্বপ্রস্তুতিমূলক স্প্রে"),
    ("General", "সাধারণ কৃষি পরামর্শ", "চাষ পদ্ধতি, বীজ বপন, জাত নির্বাচন, সেচ, জমি তৈরি, ফসল তোলার সময়"),
    ("Crisis", "জরুরি বিষক্রিয়া / নিষিদ্ধ রাসায়নিক", "কীটনাশক খেয়ে ফেলা, চোখে/মুখে বিষ লাগা, প্যারাকোয়াট, ফিউরাডন"),
    ("Off-Topic", "অকৃষি প্রশ্ন / অপ্রাসঙ্গিক", "রাজনীতি, বিনোদন, সাধারণ প্রশ্ন, আবোল-তাবোল")
]

WORKED_EXAMPLES = [
    {
        "item": "EX-01",
        "id": "sample_q_01",
        "source": "fb_group_real_farmer",
        "query": "ধানের পাতায় বাদামি দাগ দেখা যাচ্ছে এবং পাতা শুকিয়ে যাচ্ছে। এখন কি ওষুধ স্প্রে করব?",
        "crop_bn": "ধান",
        "crop_en": "rice",
        "specified": "YES",
        "intent": "Treatment",
        "confidence": "High",
        "notes": "স্পষ্টভাবে ধানের নাম এবং ব্লাইট রোগের প্রতিকার চাওয়া হয়েছে।"
    },
    {
        "item": "EX-02",
        "id": "sample_q_02",
        "source": "field_sourced_farmer",
        "query": "গাছের ডাল পচে যাচ্ছে এবং পাতা হলুদ হচ্ছে, ভালো একটা ছত্রাকনাশকের নাম বলেন।",
        "crop_bn": "কোনো ফসল নেই",
        "crop_en": "none",
        "specified": "NO",
        "intent": "Treatment",
        "confidence": "High",
        "notes": "অস্পষ্ট/ফসলহীন: গাছের উল্লেখ আছে কিন্তু কিসের গাছ (আলু/টমেটো/বেগুন) বলা নেই। রাসায়নিক পরামর্শ দেওয়া বিপজ্জনক।"
    },
    {
        "item": "EX-03",
        "id": "sample_q_03",
        "source": "krishibangla.com",
        "query": "আলু লাগানোর কত দিন পর প্রথম ইউরিয়া ও পটাশ সার উপরিপ্রয়োগ করতে হয়?",
        "crop_bn": "আলু",
        "crop_en": "potato",
        "specified": "YES",
        "intent": "Fertilizer",
        "confidence": "High",
        "notes": "নির্দিষ্ট ফসল (আলু) এবং সার প্রয়োগের সময়সূচী জানতে চাওয়া হয়েছে।"
    },
    {
        "item": "EX-04",
        "id": "sample_q_04",
        "source": "fb_group_real_farmer",
        "query": "মাজরা পোকার আক্রমণে কুশি শুকিয়ে যাচ্ছে, কি বিষ দিলে ভালো ফল পাব?",
        "crop_bn": "কোনো ফসল নেই",
        "crop_en": "none",
        "specified": "NO",
        "intent": "Treatment",
        "confidence": "High",
        "notes": "লক্ষণভিত্তিক কিন্তু ফসল অনুপস্থিত: মাজরা পোকা মূলত ধানে হলেও ভুট্টা বা আখ গাছেও হয়। ফসল নিশ্চিত না হয়ে কীটনাশক দেওয়া অনিরাপদ।"
    },
    {
        "item": "EX-05",
        "id": "sample_q_05",
        "source": "field_sourced_farmer",
        "query": "স্প্রে করার সময় মুখে কীটনাশক ঢুকে বমি ও মাথা ঘোরা শুরু হয়েছে, এখন কি করণীয়?",
        "crop_bn": "কোনো ফসল নেই",
        "crop_en": "none",
        "specified": "NO",
        "intent": "Crisis",
        "confidence": "High",
        "notes": "জরুরি বিষক্রিয়ার ঘটনা: তাৎক্ষণিক ১৬১২৩ / ৯৯৯ বা হাসপাতালে রেফার করতে হবে।"
    }
]

# Styling definitions
header_fill_context = PatternFill(start_color="1B4332", end_color="1B4332", fill_type="solid") # Dark Forest Green
header_fill_input = PatternFill(start_color="2D6A4F", end_color="2D6A4F", fill_type="solid") # Emerald Green
header_fill_ref = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid") # Deep Navy
section_fill = PatternFill(start_color="E2E8F0", end_color="E2E8F0", fill_type="solid") # Soft Gray Header

header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
sub_header_font = Font(name="Segoe UI", size=11, bold=True, color="1E293B")
font_regular = Font(name="Segoe UI", size=10, color="0F172A")
font_id = Font(name="Consolas", size=10, color="475569")
font_query = Font(name="Segoe UI", size=10.5, color="0F172A")
font_bold = Font(name="Segoe UI", size=10, bold=True, color="0F172A")

row_even_fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
row_odd_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
example_fill = PatternFill(start_color="F0FDF4", end_color="F0FDF4", fill_type="solid") # Very light mint for examples

thin_border = Border(
    left=Side(style='thin', color='E2E8F0'),
    right=Side(style='thin', color='E2E8F0'),
    top=Side(style='thin', color='E2E8F0'),
    bottom=Side(style='thin', color='E2E8F0')
)

headers = [
    ("Item #", 8, "center", header_fill_context),
    ("Query ID", 16, "center", header_fill_context),
    ("Source Channel", 22, "left", header_fill_context),
    ("Farmer Bengali Query (বাংলা প্রশ্ন)", 62, "left", header_fill_context),
    ("Crop Name (বাংলায় ফসলের নাম)", 26, "left", header_fill_input),
    ("Crop Taxonomy (English) [Dropdown]", 26, "left", header_fill_input),
    ("Is Crop Specified? (ফসল আছে?)", 22, "center", header_fill_input),
    ("Intent Category (উদ্দেশ্য)", 26, "left", header_fill_input),
    ("Confidence (আত্মবিশ্বাস)", 18, "center", header_fill_input),
    ("Reviewer Notes (মন্তব্য)", 35, "left", header_fill_input)
]

def add_reference_sheet(wb):
    ws_ref = wb.create_sheet(title="Reference_and_Examples")
    ws_ref.views.sheetView[0].showGridLines = True

    # 1. Title
    ws_ref.cell(row=1, column=1, value="CONTROLLED TAXONOMY & WORKED EXAMPLES (রেফারেন্স ও নমুনা উদাহরণ)").font = Font(name="Segoe UI", size=13, bold=True, color="1E3A8A")
    ws_ref.row_dimensions[1].height = 28

    # 2. Section 1: Crops
    ws_ref.cell(row=3, column=1, value="TABLE 1: CONTROLLED CROP VOCABULARY (ফসলের নিয়ন্ত্রিত তালিকা)").font = sub_header_font
    crop_headers = ["Crop Code (English Dropdown)", "Bangla Name (বাংলা নাম)", "Botanical Family", "Common Aliases / Variety Names (আঞ্চলিক নাম ও জাত)"]
    ws_ref.row_dimensions[4].height = 24
    for c_idx, h in enumerate(crop_headers, start=1):
        cell = ws_ref.cell(row=4, column=c_idx, value=h)
        cell.fill = header_fill_ref
        cell.font = header_font
        cell.alignment = Alignment(horizontal="left", vertical="center")
        cell.border = thin_border

    ws_ref.column_dimensions["A"].width = 28
    ws_ref.column_dimensions["B"].width = 26
    ws_ref.column_dimensions["C"].width = 22
    ws_ref.column_dimensions["D"].width = 45

    for r_offset, (code, bn, fam, aliases) in enumerate(CROPS_TAXONOMY, start=5):
        fill = row_even_fill if r_offset % 2 == 0 else row_odd_fill
        ws_ref.row_dimensions[r_offset].height = 20
        c1 = ws_ref.cell(row=r_offset, column=1, value=code)
        c2 = ws_ref.cell(row=r_offset, column=2, value=bn)
        c3 = ws_ref.cell(row=r_offset, column=3, value=fam)
        c4 = ws_ref.cell(row=r_offset, column=4, value=aliases)

        c1.font = font_bold if code in ["none", "rice", "potato", "wheat"] else font_regular
        c2.font = font_regular
        c3.font = font_id
        c4.font = font_regular

        for c in [c1, c2, c3, c4]:
            c.fill = fill
            c.border = thin_border
            c.alignment = Alignment(horizontal="left", vertical="center")

    last_crop_row = 4 + len(CROPS_TAXONOMY)

    # 3. Section 2: Intent Options
    intent_start_row = last_crop_row + 3
    ws_ref.cell(row=intent_start_row, column=1, value="TABLE 2: INTENT CATEGORY TAXONOMY (প্রশ্নের ধরণ)").font = sub_header_font
    ws_ref.row_dimensions[intent_start_row+1].height = 24
    intent_headers = ["Intent Code (Dropdown)", "Bangla Description (বাংলা অর্থ)", "Typical Scope & Trigger Keywords (লক্ষণ ও মূল শব্দ)"]
    for c_idx, h in enumerate(intent_headers, start=1):
        cell = ws_ref.cell(row=intent_start_row+1, column=c_idx, value=h)
        cell.fill = header_fill_ref
        cell.font = header_font
        cell.alignment = Alignment(horizontal="left", vertical="center")
        cell.border = thin_border

    for r_offset, (icode, ibn, idesc) in enumerate(INTENT_OPTIONS, start=intent_start_row+2):
        fill = row_even_fill if r_offset % 2 == 0 else row_odd_fill
        ws_ref.row_dimensions[r_offset].height = 24
        c1 = ws_ref.cell(row=r_offset, column=1, value=icode)
        c2 = ws_ref.cell(row=r_offset, column=2, value=ibn)
        c3 = ws_ref.cell(row=r_offset, column=3, value=idesc)
        for c in [c1, c2, c3]:
            c.fill = fill
            c.font = font_bold if c == c1 else font_regular
            c.border = thin_border
            c.alignment = Alignment(horizontal="left", vertical="center")

    last_intent_row = intent_start_row + 1 + len(INTENT_OPTIONS)

    # 4. Section 3: Worked Examples
    ex_start_row = last_intent_row + 3
    ws_ref.cell(row=ex_start_row, column=1, value="TABLE 3: WORKED EXAMPLES (নমুনা পূর্ণাঙ্গ এ্যানোটেশন)").font = sub_header_font
    ws_ref.row_dimensions[ex_start_row+1].height = 26
    ex_cols = [
        ("Item #", 10), ("ID", 14), ("Query (বাংলা প্রশ্ন)", 55), ("Crop (বাংলা)", 18),
        ("Crop (English)", 18), ("Specified?", 14), ("Intent", 16), ("Notes / Explanation (মন্তব্য ও যুক্তি)", 45)
    ]
    for c_idx, (eh, ew) in enumerate(ex_cols, start=1):
        cell = ws_ref.cell(row=ex_start_row+1, column=c_idx, value=eh)
        cell.fill = header_fill_context
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center" if c_idx in [1,2,6] else "left", vertical="center")
        cell.border = thin_border

    for r_offset, ex in enumerate(WORKED_EXAMPLES, start=ex_start_row+2):
        ws_ref.row_dimensions[r_offset].height = 42
        row_vals = [
            ex["item"], ex["id"], ex["query"], ex["crop_bn"],
            ex["crop_en"], ex["specified"], ex["intent"], ex["notes"]
        ]
        for c_idx, val in enumerate(row_vals, start=1):
            cell = ws_ref.cell(row=r_offset, column=c_idx, value=val)
            cell.fill = example_fill
            cell.font = font_bold if c_idx in [4,5,6] else font_regular
            cell.border = thin_border
            cell.alignment = Alignment(horizontal="center" if c_idx in [1,2,6] else "left", vertical="center", wrap_text=(c_idx in [3,8]))

    return last_crop_row

def style_worksheet(ws, items, start_num=1, last_crop_row=41):
    ws.views.sheetView[0].showGridLines = True
    ws.freeze_panes = "E2"
    
    # Header
    ws.row_dimensions[1].height = 32
    for col_idx, (title, width, align, fill) in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx, value=title)
        cell.font = header_font
        cell.fill = fill
        cell.alignment = Alignment(horizontal=align, vertical="center", wrap_text=True)
        cell.border = thin_border
        col_letter = get_column_letter(col_idx)
        ws.column_dimensions[col_letter].width = width

    # Data Validations
    # Crop Taxonomy dropdown references Reference_and_Examples sheet!
    dv_crop_taxonomy = DataValidation(type="list", formula1=f"Reference_and_Examples!$A$5:$A${last_crop_row}", allow_blank=True)
    dv_crop_specified = DataValidation(type="list", formula1='"YES,NO"', allow_blank=True)
    dv_intent = DataValidation(type="list", formula1='"Treatment,Fertilizer,Prevention,General,Crisis,Off-Topic"', allow_blank=True)
    dv_confidence = DataValidation(type="list", formula1='"High,Medium,Low"', allow_blank=True)

    ws.add_data_validation(dv_crop_taxonomy)
    ws.add_data_validation(dv_crop_specified)
    ws.add_data_validation(dv_intent)
    ws.add_data_validation(dv_confidence)

    # Populate rows
    for r_idx, q in enumerate(items, start=2):
        item_num = start_num + r_idx - 2
        ws.row_dimensions[r_idx].height = 54
        fill = row_even_fill if r_idx % 2 == 0 else row_odd_fill
        
        ws.cell(row=r_idx, column=1, value=item_num)
        ws.cell(row=r_idx, column=2, value=q.get("id"))
        ws.cell(row=r_idx, column=3, value=q.get("source"))
        ws.cell(row=r_idx, column=4, value=q.get("query"))
        
        for c in range(5, 11):
            ws.cell(row=r_idx, column=c, value="")

        ws.cell(row=r_idx, column=1).alignment = Alignment(horizontal="center", vertical="center")
        ws.cell(row=r_idx, column=1).font = font_regular

        ws.cell(row=r_idx, column=2).alignment = Alignment(horizontal="center", vertical="center")
        ws.cell(row=r_idx, column=2).font = font_id

        ws.cell(row=r_idx, column=3).alignment = Alignment(horizontal="left", vertical="center")
        ws.cell(row=r_idx, column=3).font = font_regular

        ws.cell(row=r_idx, column=4).alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        ws.cell(row=r_idx, column=4).font = font_query

        ws.cell(row=r_idx, column=5).alignment = Alignment(horizontal="left", vertical="center")
        ws.cell(row=r_idx, column=5).font = font_regular

        # Column 6: Crop Taxonomy Dropdown
        c6 = ws.cell(row=r_idx, column=6)
        c6.alignment = Alignment(horizontal="left", vertical="center")
        c6.font = font_regular
        dv_crop_taxonomy.add(c6)

        # Column 7: Specified Dropdown
        c7 = ws.cell(row=r_idx, column=7)
        c7.alignment = Alignment(horizontal="center", vertical="center")
        c7.font = font_regular
        dv_crop_specified.add(c7)

        # Column 8: Intent Dropdown
        c8 = ws.cell(row=r_idx, column=8)
        c8.alignment = Alignment(horizontal="left", vertical="center")
        c8.font = font_regular
        dv_intent.add(c8)

        # Column 9: Confidence Dropdown
        c9 = ws.cell(row=r_idx, column=9)
        c9.alignment = Alignment(horizontal="center", vertical="center")
        c9.font = font_regular
        dv_confidence.add(c9)

        # Column 10: Notes
        ws.cell(row=r_idx, column=10).alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        ws.cell(row=r_idx, column=10).font = font_regular

        for c_idx in range(1, 11):
            cell = ws.cell(row=r_idx, column=c_idx)
            cell.fill = fill
            cell.border = thin_border

# Generate 10 batches of 20 queries each
BATCH_SIZE = 20
total_batches = (len(queries) + BATCH_SIZE - 1) // BATCH_SIZE

for b_idx in range(total_batches):
    start = b_idx * BATCH_SIZE
    end = min(start + BATCH_SIZE, len(queries))
    batch_items = queries[start:end]
    
    wb = openpyxl.Workbook()
    # First add reference sheet
    last_crop_row = add_reference_sheet(wb)
    
    # Active annotation sheet
    ws = wb.worksheets[0]
    ws.title = f"Batch_{b_idx+1:02d}"
    
    style_worksheet(ws, batch_items, start_num=start+1, last_crop_row=last_crop_row)
    
    # Put annotation sheet first
    wb.active = ws
    
    fname = f"batch_{b_idx+1:02d}_queries_{start+1:03d}_{end:03d}.xlsx"
    fpath = os.path.join(BASE_DIR, "batches", fname)
    wb.save(fpath)
    print(f"Generated {fname} ({len(batch_items)} rows + Reference sheet)")

# Generate Master Workbook containing all 200 queries
wb_master = openpyxl.Workbook()
last_crop_row = add_reference_sheet(wb_master)
ws_all = wb_master.worksheets[0]
ws_all.title = "All 200 Queries"
style_worksheet(ws_all, queries, start_num=1, last_crop_row=last_crop_row)

for b_idx in range(total_batches):
    start = b_idx * BATCH_SIZE
    end = min(start + BATCH_SIZE, len(queries))
    batch_items = queries[start:end]
    ws_tab = wb_master.create_sheet(title=f"Batch {b_idx+1:02d}")
    style_worksheet(ws_tab, batch_items, start_num=start+1, last_crop_row=last_crop_row)

wb_master.active = ws_all
master_path = os.path.join(BASE_DIR, "master_blind_sheet_200.xlsx")
try:
    wb_master.save(master_path)
    print(f"Generated master workbook: {master_path}")
except PermissionError:
    alt_path = os.path.join(BASE_DIR, "master_blind_sheet_200_with_dropdowns.xlsx")
    wb_master.save(alt_path)
    print(f"master_blind_sheet_200.xlsx is currently open in Excel. Saved updated version with dropdowns to: {alt_path}")
