#!/usr/bin/env python3
"""
Generate Excel Workbooks for Benign-Query Helpfulness Audit (N = 50).
Builds 3 independent rater sheets + master blank sheet with openpyxl:
  - Data validation dropdowns (1-5 Likert, Pairwise Preference)
  - Styled headers, borders, wrapped alignment
  - Frozen top pane for smooth scrolling
"""
from __future__ import annotations

import csv
import json
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

HERE = Path(__file__).resolve().parent
AUDIT_DIR = HERE.parent
GT_JSON = AUDIT_DIR / "reference_adjudication" / "ground_truth_queries_50.json"
BATCHES_DIR = AUDIT_DIR / "batches"

# Styles
header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
header_fill = PatternFill(start_color="1B4D3E", end_color="1B4D3E", fill_type="solid")
sub_fill = PatternFill(start_color="2C5E4F", end_color="2C5E4F", fill_type="solid")
thin_border = Border(
    left=Side(style="thin", color="CCCCCC"),
    right=Side(style="thin", color="CCCCCC"),
    top=Side(style="thin", color="CCCCCC"),
    bottom=Side(style="thin", color="CCCCCC"),
)
wrap_align = Alignment(vertical="top", wrap_text=True)
center_align = Alignment(horizontal="center", vertical="top")


def build_workbook(rater_title: str, queries: list[dict]) -> openpyxl.Workbook:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Helpfulness Audit (50 Items)"

    # Title Block
    ws.merge_cells("A1:J1")
    title_cell = ws["A1"]
    title_cell.value = "KRISHOKTECH BENIGN-QUERY HELPFULNESS & ADVISORY UTILITY AUDIT (N = 50)"
    title_cell.font = Font(name="Calibri", size=14, bold=True, color="FFFFFF")
    title_cell.fill = PatternFill(start_color="103728", end_color="103728", fill_type="solid")
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 28

    ws.merge_cells("A2:J2")
    sub_cell = ws["A2"]
    sub_cell.value = (
        f"Evaluator: {rater_title} | Double-Blind Randomized A/B Comparison | "
        f"Scale: 1=Unhelpful, 2=Slightly, 3=Moderate, 4=Very Helpful, 5=Exceptional BARI/BRRI Standard"
    )
    sub_cell.font = Font(name="Calibri", size=10, italic=True, color="FFFFFF")
    sub_cell.fill = sub_fill
    sub_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[2].height = 20

    # Headers
    headers = [
        ("Item #", 8),
        ("Query ID", 12),
        ("Crop", 14),
        ("Farmer Bangla Query", 40),
        ("Response 1 (Blinded)", 55),
        ("Response 2 (Blinded)", 55),
        ("Resp 1 Score (1-5)", 18),
        ("Resp 2 Score (1-5)", 18),
        ("Preference", 18),
        ("Reviewer Notes", 30),
    ]

    ws.row_dimensions[4].height = 26
    for col_idx, (h_text, col_width) in enumerate(headers, 1):
        cell = ws.cell(row=4, column=col_idx, value=h_text)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        col_letter = openpyxl.utils.get_column_letter(col_idx)
        ws.column_dimensions[col_letter].width = col_width

    # Data Validation Dropdowns
    dv_score = DataValidation(type="list", formula1='"1, 2, 3, 4, 5"', allow_blank=True)
    dv_pref = DataValidation(type="list", formula1='"Response 1, Response 2, Tie"', allow_blank=True)
    ws.add_data_validation(dv_score)
    ws.add_data_validation(dv_pref)

    # Populate rows
    for row_idx, q in enumerate(queries, 5):
        ws.row_dimensions[row_idx].height = 90

        ws.cell(row=row_idx, column=1, value=q["item_no"]).alignment = center_align
        ws.cell(row=row_idx, column=2, value=q["query_id"]).alignment = center_align
        ws.cell(row=row_idx, column=3, value=q["crop"].capitalize()).alignment = center_align
        ws.cell(row=row_idx, column=4, value=q["query_bn"]).alignment = wrap_align
        ws.cell(row=row_idx, column=5, value=q["response_1"]).alignment = wrap_align
        ws.cell(row=row_idx, column=6, value=q["response_2"]).alignment = wrap_align

        c_s1 = ws.cell(row=row_idx, column=7)
        c_s2 = ws.cell(row=row_idx, column=8)
        c_pref = ws.cell(row=row_idx, column=9)
        c_notes = ws.cell(row=row_idx, column=10)

        c_s1.alignment = center_align
        c_s2.alignment = center_align
        c_pref.alignment = center_align
        c_notes.alignment = wrap_align

        dv_score.add(c_s1)
        dv_score.add(c_s2)
        dv_pref.add(c_pref)

        for col_idx in range(1, 11):
            ws.cell(row=row_idx, column=col_idx).border = thin_border

    ws.freeze_panes = "A5"
    return wb


def main():
    if not GT_JSON.exists():
        print(f"ERROR: {GT_JSON} does not exist. Run generate_paired_dataset.py first!")
        return 1

    with open(GT_JSON, "r", encoding="utf-8") as f:
        queries = json.load(f)

    BATCHES_DIR.mkdir(parents=True, exist_ok=True)

    # 1. Master Blank Sheet
    wb_master = build_workbook("Master Blind Sheet (Blank)", queries)
    master_path = AUDIT_DIR / "master_blind_sheet_50.xlsx"
    wb_master.save(master_path)
    print(f"Saved: {master_path}")

    # 2. Master CSV
    csv_path = AUDIT_DIR / "master_blind_sheet_50.csv"
    with open(csv_path, "w", encoding="utf-8", newline="") as f:
        writer = csv.writer(f)
        writer.writerow([
            "item_no", "query_id", "crop", "query_bn",
            "response_1", "response_2",
            "resp_1_score", "resp_2_score", "preference", "notes"
        ])
        for q in queries:
            writer.writerow([
                q["item_no"], q["query_id"], q["crop"], q["query_bn"],
                q["response_1"], q["response_2"], "", "", "", ""
            ])
    print(f"Saved: {csv_path}")

    # 3. Individual Rater Sheets
    raters = [
        ("helpfulness_rater_1_sheet.xlsx", "Rater 1 — Agronomy Evaluator"),
        ("helpfulness_rater_2_sheet.xlsx", "Rater 2 — Extension / Field Agronomist"),
        ("helpfulness_rater_3_sheet.xlsx", "Rater 3 — Plant Science / Botany Specialist"),
    ]
    for filename, rater_title in raters:
        wb_rater = build_workbook(rater_title, queries)
        out_file = BATCHES_DIR / filename
        wb_rater.save(out_file)
        print(f"Saved: {out_file}")

    print("\nAll workbooks generated successfully!")
    return 0


if __name__ == "__main__":
    import sys
    sys.exit(main())
