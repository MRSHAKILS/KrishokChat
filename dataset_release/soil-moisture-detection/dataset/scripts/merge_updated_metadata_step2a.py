"""
merge_updated_metadata_step2a.py
Merge TraceData.xlsx metadata + correct_image/ into the canonical dataset.
Outputs: image_manifest.csv with all columns (initially v3, later finalized by step2d).
"""
import csv, json, sys
from pathlib import Path
from datetime import time
from collections import defaultdict, Counter

try:
    import openpyxl
except ImportError:
    print("ERROR: openpyxl required. pip install openpyxl")
    sys.exit(1)

SCRIPT_DIR = Path(__file__).resolve().parent
DATASET_DIR = SCRIPT_DIR.parent
PREPROC_DIR = DATASET_DIR / "preprocessed"
UPDATED_DIR = DATASET_DIR / "updated_dataset"
CLEAN_DIR = DATASET_DIR / "clean"

# ============================================================
# 1. Load TraceData.xlsx
# ============================================================
xlsx_path = UPDATED_DIR / "TraceData.xlsx"
if not xlsx_path.exists():
    print(f"ERROR: {xlsx_path} not found")
    sys.exit(1)

wb = openpyxl.load_workbook(xlsx_path)
ws = wb["Sheet1"]

# Read header
headers = [cell.value for cell in ws[1]]
print(f"TraceData columns: {headers}")
# ['ID', 'FILE_NAME', 'Kpa', 'LAND_TYPE', 'SOIL_TYPE',
#  'LAST_IRREG/RAIN_DATE', 'LAST_IRREG/RAIN_TIME',
#  'CROPS_NAME', 'GROWTH_STAGE', 'DATE', 'TIME', 'Locatio']

trace_data = {}  # filename -> row
for row in ws.iter_rows(min_row=2, values_only=True):
    id_val, fname, kpa, land_type, soil_type, rain_date, rain_time, crop, growth_stage, date_val, time_val, loc = row
    if fname is None:
        continue
    # Normalize filename
    fname = str(fname).strip()
    trace_data[fname] = {
        "file_name": fname,
        "kpa": str(kpa) if kpa is not None else "",
        "land_type": str(land_type) if land_type else "",
        "soil_type": str(soil_type) if soil_type else "",
        "rain_date": str(rain_date) if rain_date else "",
        "rain_time": str(rain_time) if rain_time else "",
        "crop": str(crop) if crop else "",
        "growth_stage": str(growth_stage) if growth_stage else "",
        "date": str(date_val) if date_val else "",
        "time": str(time_val) if time_val else "",
        "location": str(loc) if loc else "",
    }

print(f"TraceData entries: {len(trace_data)}")

# Stats
soil_counts = Counter(d["soil_type"] for d in trace_data.values())
land_counts = Counter(d["land_type"] for d in trace_data.values())
crop_counts = Counter(d["crop"] for d in trace_data.values())
stage_counts = Counter(d["growth_stage"] for d in trace_data.values())

print(f"\nSOIL_TYPE distribution:")
for k, v in sorted(soil_counts.items(), key=lambda x: -x[1]):
    print(f"  {k}: {v}")
print(f"\nLAND_TYPE distribution:")
for k, v in sorted(land_counts.items(), key=lambda x: -x[1]):
    print(f"  {k}: {v}")
print(f"\nCROPS distribution (top 10):")
for k, v in sorted(crop_counts.items(), key=lambda x: -x[1])[:10]:
    print(f"  {k}: {v}")
print(f"\nGROWTH_STAGE distribution:")
for k, v in sorted(stage_counts.items(), key=lambda x: -x[1]):
    print(f"  {k}: {v}")

# ============================================================
# 2. Compare correct_image/ vs clean/
# ============================================================
correct_dir = UPDATED_DIR / "correct_image"
correct_files = {f.name: f for f in correct_dir.glob("*.jpg")}
clean_files = {f.name: f for f in CLEAN_DIR.glob("*.jpg")}

print(f"\ncorrect_image/ files: {len(correct_files)}")
print(f"clean/ files: {len(clean_files)}")

# Files in correct but not in clean
only_correct = set(correct_files.keys()) - set(clean_files.keys())
only_clean = set(clean_files.keys()) - set(correct_files.keys())
common = set(correct_files.keys()) & set(clean_files.keys())

if only_correct:
    print(f"\nFiles ONLY in correct_image/ (not in clean/): {len(only_correct)}")
    for f in sorted(only_correct)[:10]:
        print(f"  {f}")
if only_clean:
    print(f"\nFiles ONLY in clean/ (not in correct_image/): {len(only_clean)}")
    for f in sorted(only_clean)[:10]:
        print(f"  {f}")

# Compare file sizes for common files (indicates content differences)
size_diffs = []
for fname in common:
    s1 = correct_files[fname].stat().st_size
    s2 = clean_files[fname].stat().st_size
    if s1 != s2:
        size_diffs.append((fname, s1, s2, abs(s1-s2)))

size_diffs.sort(key=lambda x: -x[3])

if size_diffs:
    print(f"\nFiles with DIFFERENT size between correct/ and clean/ (top 20):")
    for fname, s1, s2, diff in size_diffs[:20]:
        print(f"  {fname}: correct={s1} vs clean={s2} (diff={diff})")
    print(f"  Total different: {len(size_diffs)} / {len(common)}")
else:
    print(f"\nAll {len(common)} common files have identical sizes — correct/ == clean/")

# ============================================================
# 3. Merge with existing manifest
# ============================================================
print(f"\n{'='*60}")
print("MERGING METADATA INTO MANIFEST")
print(f"{'='*60}")

# Load existing manifest
manifest_rows = []
with open(PREPROC_DIR / "image_manifest.csv", "r", encoding="utf-8") as f:
    reader = csv.DictReader(f)
    for row in reader:
        manifest_rows.append(row)

print(f"Existing manifest: {len(manifest_rows)} rows")

# Match by filename — handle clean filenames vs raw filenames
# TraceData uses raw filenames (P0001_8Kpa.jpg), manifest uses clean (P0001_8.0Kpa.jpg)
# Build mapping from raw -> clean using rename_log
rename_map = {}
with open(PREPROC_DIR / "rename_log.csv", "r", encoding="utf-8") as f:
    reader = csv.DictReader(f)
    for r in reader:
        rename_map[r["original"]] = r["clean"]

# Reverse: clean -> raw
clean_to_raw = {v: k for k, v in rename_map.items()}

# Build new manifest with merged metadata
merged = []
match_count = 0
missing_metadata = []
new_fields = [
    "land_type", "soil_type", "rain_date", "rain_time",
    "crop", "growth_stage", "capture_date", "capture_time", "location"
]

for row in manifest_rows:
    clean_fn = row["filename"]
    raw_fn = clean_to_raw.get(clean_fn, clean_fn)
    
    meta = trace_data.get(raw_fn, {})
    if not meta:
        # Try direct match
        meta = trace_data.get(clean_fn, {})
    
    if meta:
        match_count += 1
    else:
        missing_metadata.append(clean_fn)
    
    merged.append({
        "image_id": row["image_id"],
        "original_id": row["original_id"],
        "filename": clean_fn,
        "kpa_raw": row["kpa_raw"],
        "kpa_bin": row["kpa_bin"],
        "anomalies": row["anomalies"],
        "id_num": row["id_num"],
        "series_id": row["series_id"],
        "series_size": row["series_size"],
        "land_type": meta.get("land_type", ""),
        "soil_type": meta.get("soil_type", ""),
        "rain_date": meta.get("rain_date", ""),
        "rain_time": meta.get("rain_time", ""),
        "crop": meta.get("crop", ""),
        "growth_stage": meta.get("growth_stage", ""),
        "capture_date": meta.get("date", ""),
        "capture_time": meta.get("time", ""),
        "location": meta.get("location", ""),
    })

print(f"Matched: {match_count}/{len(manifest_rows)}")
if missing_metadata:
    print(f"Missing metadata: {len(missing_metadata)}")
    for f in missing_metadata[:10]:
        print(f"  {f}")

# ============================================================
# 4. Write merged manifest
# ============================================================
fieldnames = [
    "image_id", "original_id", "filename", "kpa_raw", "kpa_bin",
    "anomalies", "id_num", "series_id", "series_size",
    "land_type", "soil_type", "rain_date", "rain_time",
    "crop", "growth_stage", "capture_date", "capture_time", "location"
]

output_path = PREPROC_DIR / "image_manifest.csv"
with open(output_path, "w", encoding="utf-8", newline="") as f:
    writer = csv.DictWriter(f, fieldnames=fieldnames, quoting=csv.QUOTE_ALL)
    writer.writeheader()
    for row in merged:
        writer.writerow(row)

print(f"\nWritten: {output_path} ({len(merged)} rows)")

# ============================================================
# 5. Report: cross-tab soil_type x kpa_bin
# ============================================================
print(f"\n{'='*60}")
print("SOIL TYPE x kPa BIN CROSS-TABULATION")
print(f"{'='*60}")

cross = defaultdict(lambda: defaultdict(int))
soil_bins = defaultdict(int)
bin_soils = defaultdict(int)
for row in merged:
    st = row["soil_type"] or "UNKNOWN"
    kb = row["kpa_bin"] or "UNKNOWN"
    cross[st][kb] += 1
    soil_bins[st] += 1
    bin_soils[kb] += 1

soil_types = sorted(soil_bins.keys(), key=lambda x: -soil_bins[x])
kpa_bins = ["0-10", "10-20", "20-30", ">100"]

header = f"{'SOIL TYPE':<20s}" + "".join(f"{b:>8s}" for b in kpa_bins) + f"{'TOTAL':>8s}"
print(header)
print("-" * len(header))
for st in soil_types:
    vals = [str(cross[st][b]) for b in kpa_bins]
    print(f"{st:<20s}" + "".join(f"{v:>8s}" for v in vals) + f"{soil_bins[st]:>8d}")

# ============================================================
# 6. Report: crop x soil_type
# ============================================================
print(f"\n{'='*60}")
print("CROP x SOIL TYPE")
print(f"{'='*60}")
crop_soil = defaultdict(lambda: defaultdict(int))
for row in merged:
    c = row["crop"] or "UNKNOWN"
    s = row["soil_type"] or "UNKNOWN"
    crop_soil[c][s] += 1

for crop in sorted(crop_soil.keys()):
    soils = crop_soil[crop]
    total = sum(soils.values())
    soil_str = ", ".join(f"{k}={v}" for k, v in sorted(soils.items(), key=lambda x: -x[1]))
    print(f"  {crop:<20s} ({total:3d}): {soil_str}")

# ============================================================
# 7. Date/time range
# ============================================================
dates = set()
times = set()
for row in merged:
    if row["capture_date"] and row["capture_date"] != "None":
        dates.add(row["capture_date"])
    if row["capture_time"] and row["capture_time"] != "None":
        times.add(row["capture_time"])

print(f"\nCapture dates: {sorted(dates) if dates else 'Not available'}")
print(f"Capture times: {sorted(times)[:10] if times else 'Not available'}...")

rain_dates = set()
for row in merged:
    if row["rain_date"] and row["rain_date"] != "None":
        rain_dates.add(row["rain_date"])
print(f"Last irrigation dates: {sorted(rain_dates) if rain_dates else 'Not available'}")

print(f"\nDone. Merged metadata written to image_manifest.csv")
