"""
Metadata Sheet Generator
=========================
Creates the Excel annotation template for dataset labeling.

Run from: dataset/scripts/
Output:   dataset/metadata/metadata.xlsx
"""

import os
import re
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from pathlib import Path

# ── Path Configuration ─────────────────────────────────────────────────────
SCRIPT_DIR = Path(__file__).resolve().parent
DATASET_DIR = SCRIPT_DIR.parent
RAW_DIR = DATASET_DIR / "raw"
OUTPUT_PATH = DATASET_DIR / "metadata" / "metadata.xlsx"


def get_kpa_range(filename):
    """Parse kPa from filename and return bin label."""
    parts = filename.split("_")
    if len(parts) < 2:
        return ""
    match = re.search(r"([0-9.,]+)", parts[1])
    if not match:
        return ""
    num_str = match.group(1).replace(",", ".")
    if num_str == "985":
        val = 9.85
    else:
        try:
            val = float(num_str)
        except ValueError:
            return ""
    if 0 <= val <= 10:
        return "0-10"
    elif 10 < val <= 20:
        return "10-20"
    elif 20 < val <= 30:
        return "20-30"
    elif 30 < val <= 40:
        return "30-40"
    elif 40 < val <= 50:
        return "40-50"
    elif 50 < val <= 60:
        return "50-60"
    elif 60 < val <= 70:
        return "60-70"
    elif 70 < val <= 80:
        return "70-80"
    elif 80 < val <= 90:
        return "80-90"
    elif 90 < val <= 100:
        return "90-100"
    return ""


def main():
    if not RAW_DIR.exists():
        print(f"Error: Raw images not found at {RAW_DIR}")
        return

    extensions = (".jpg", ".jpeg", ".png")
    image_files = sorted([f for f in os.listdir(RAW_DIR) if f.lower().endswith(extensions)])

    if not image_files:
        print("No image files found!")
        return
    print(f"Found {len(image_files)} images. Generating sheet...")

    dropdown_options = {
        "soil_type": ["Loamy", "Sandy-Loamy", "Clay", "Clay-Loamy", "Sandy"],
        "crop_name": [
            "Boro Rice", "Aman Rice", "Aus Rice", "Wheat", "Potato",
            "Onion", "Garlic", "Mustard", "Corn", "Tomato",
            "Brinjal", "Chilli", "Cauliflower", "Cabbage", "Bottle Gourd",
            "Pumpkin", "Jute", "Sugarcane", "Lentil", "Chickpea"
        ],
        "kPa": ["0-10", "10-20", "20-30", "30-40", "40-50",
                "50-60", "60-70", "70-80", "80-90", "90-100"],
        "temp_C": [20, 22, 24, 26, 28, 30, 32, 34, 36, 38, 40, 42],
        "land_type": ["High", "Medium", "Low"],
        "growth_stage": ["Seed Sowing", "Seedling", "Tillering", "Flowering", "Fruiting", "Maturity"],
        "weather": ["Sunny", "Partly Cloudy", "Cloudy", "After Rain", "Before Rain"],
        "last_irrigated": ["Today", "1 Day Ago", "2 Days Ago", "3 Days Ago", "4 Days Ago", "5 Days Ago", "1 Week Ago", "Not Irrigated"],
        "location": ["Ishwardi", "Pabna Sadar", "Rajapur", "Shibpur", "Atgharia", "Bera", "Sujanagar"],
        "time": ["6:00 AM", "7:00 AM", "8:00 AM", "9:00 AM", "12:00 PM", "3:00 PM", "5:00 PM"]
    }

    wb = openpyxl.Workbook()
    ws_meta = wb.active
    ws_meta.title = "Metadata"
    ws_lists = wb.create_sheet(title="List_Options")

    col_mapping = {}
    for idx, (key, values) in enumerate(dropdown_options.items(), 1):
        col_letter = get_column_letter(idx)
        cell_header = ws_lists.cell(row=1, column=idx, value=f"{key}_options")
        cell_header.font = Font(name="Segoe UI", bold=True, color="5B5D5F")
        cell_header.fill = PatternFill(start_color="EBECEC", end_color="EBECEC", fill_type="solid")
        for r_idx, val in enumerate(values, 2):
            ws_lists.cell(row=r_idx, column=idx, value=val)
        col_mapping[key] = f"List_Options!${col_letter}$2:${col_letter}${len(values) + 1}"

    headers = [
        "image_id", "filename", "soil_type", "crop_name", "kPa", "temp_C",
        "land_type", "growth_stage", "weather", "last_irrigated", "location",
        "date", "time", "datetime"
    ]

    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    font_data = Font(name="Segoe UI", size=10)
    font_formula = Font(name="Segoe UI", size=10, italic=True)
    align_center = Alignment(horizontal="center", vertical="center")
    align_left = Alignment(horizontal="left", vertical="center")
    border_thin = Side(border_style="thin", color="D5D8DC")
    border_double = Side(border_style="double", color="1F4E78")
    cell_border = Border(left=border_thin, right=border_thin, top=border_thin, bottom=border_thin)
    header_border = Border(left=border_thin, right=border_thin, top=border_thin, bottom=border_double)
    fill_even = PatternFill(start_color="F7F9FC", end_color="F7F9FC", fill_type="solid")

    for col_idx, header in enumerate(headers, 1):
        cell = ws_meta.cell(row=1, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = header_border
    ws_meta.row_dimensions[1].height = 28

    for row_idx, fname in enumerate(image_files, 2):
        image_id = fname.split("_")[0] if "_" in fname else os.path.splitext(fname)[0]
        kpa_range = get_kpa_range(fname)
        ws_meta.cell(row=row_idx, column=1, value=image_id)
        ws_meta.cell(row=row_idx, column=2, value=fname)
        ws_meta.cell(row=row_idx, column=5, value=kpa_range)
        ws_meta.row_dimensions[row_idx].height = 20
        row_fill = fill_even if row_idx % 2 == 0 else PatternFill(fill_type=None)
        for col_idx in range(1, len(headers) + 1):
            cell = ws_meta.cell(row=row_idx, column=col_idx)
            cell.font = font_data
            cell.border = cell_border
            if row_fill.fill_type:
                cell.fill = row_fill
            if col_idx in [1, 5, 6, 7, 12, 13, 14]:
                cell.alignment = align_center
            else:
                cell.alignment = align_left
        formula_str = f'=IF(OR(ISBLANK(L{row_idx}), ISBLANK(M{row_idx})), "", TEXT(L{row_idx}, "DD/MM/YYYY") & " " & M{row_idx})'
        ws_meta.cell(row=row_idx, column=14, value=formula_str).font = font_formula

    validations = {
        "soil_type": DataValidation(type="list", formula1=col_mapping["soil_type"], allow_blank=True),
        "crop_name": DataValidation(type="list", formula1=col_mapping["crop_name"], allow_blank=True),
        "kPa": DataValidation(type="list", formula1=col_mapping["kPa"], allow_blank=True),
        "temp_C": DataValidation(type="list", formula1=col_mapping["temp_C"], allow_blank=True),
        "land_type": DataValidation(type="list", formula1=col_mapping["land_type"], allow_blank=True),
        "growth_stage": DataValidation(type="list", formula1=col_mapping["growth_stage"], allow_blank=True),
        "weather": DataValidation(type="list", formula1=col_mapping["weather"], allow_blank=True),
        "last_irrigated": DataValidation(type="list", formula1=col_mapping["last_irrigated"], allow_blank=True),
        "location": DataValidation(type="list", formula1=col_mapping["location"], allow_blank=True),
        "time": DataValidation(type="list", formula1=col_mapping["time"], allow_blank=True),
    }

    num_rows = len(image_files) + 1
    for val in validations.values():
        val.error = "Your entry is not in the list"
        val.errorTitle = "Invalid Value Selected"
        val.showErrorMessage = True

    for name, val_obj in validations.items():
        ws_meta.add_data_validation(val_obj)

    validations["soil_type"].add(f"C2:C{num_rows}")
    validations["crop_name"].add(f"D2:D{num_rows}")
    validations["kPa"].add(f"E2:E{num_rows}")
    validations["temp_C"].add(f"F2:F{num_rows}")
    validations["land_type"].add(f"G2:G{num_rows}")
    validations["growth_stage"].add(f"H2:H{num_rows}")
    validations["weather"].add(f"I2:I{num_rows}")
    validations["last_irrigated"].add(f"J2:J{num_rows}")
    validations["location"].add(f"K2:K{num_rows}")
    validations["time"].add(f"M2:M{num_rows}")

    date_val = DataValidation(type="custom", formula1="=ISNUMBER(L2)", allow_blank=True)
    date_val.error = "Please enter date in standard format"
    date_val.errorTitle = "Invalid Date Input"
    date_val.prompt = "Enter date (e.g., DD/MM/YYYY)"
    date_val.promptTitle = "Date Format"
    date_val.showInputMessage = True
    date_val.showErrorMessage = True
    ws_meta.add_data_validation(date_val)
    date_val.add(f"L2:L{num_rows}")

    for ws in [ws_meta, ws_lists]:
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val = str(cell.value or "")
                if val.startswith("="):
                    val = "DD/MM/YYYY HH:MM AM"
                if len(val) > max_len:
                    max_len = len(val)
            width = max(max_len + 4, 12)
            if width > 45:
                width = 45
            ws.column_dimensions[col_letter].width = width

    OUTPUT_PATH.parent.mkdir(exist_ok=True)
    wb.save(OUTPUT_PATH)
    print(f"Generated: {OUTPUT_PATH} ({len(image_files)} rows)")


if __name__ == "__main__":
    main()
